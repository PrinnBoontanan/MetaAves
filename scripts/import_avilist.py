#!/usr/bin/env python3
"""
Build MetaAves bird/taxonomy data from the official AviList v2025b XLSX.

Usage:
  python scripts/import_avilist.py path/to/AviList-v2025b-extended.xlsx

The importer treats AviList as the authoritative ranked taxonomy and
automatically finds the AviList extended worksheet, so minor worksheet
name/capitalization changes do not break the import.
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


# Broad phylogenetic backbone used by MetaAves. These assignments intentionally
# stop at stable/useful named clades rather than encoding every disputed deep
# Neoaves relationship.
CLADE_PATHS_BY_ORDER = {
    "Struthioniformes": ["Neornithes", "Palaeognathae"],
    "Casuariiformes": ["Neornithes", "Palaeognathae"],
    "Apterygiformes": ["Neornithes", "Palaeognathae"],
    "Rheiformes": ["Neornithes", "Palaeognathae"],
    "Tinamiformes": ["Neornithes", "Palaeognathae"],

    "Anseriformes": ["Neornithes", "Neognathae", "Galloanserae"],
    "Galliformes": ["Neornithes", "Neognathae", "Galloanserae"],

    "Phoenicopteriformes": ["Neornithes", "Neognathae", "Neoaves", "Mirandornithes"],
    "Podicipediformes": ["Neornithes", "Neognathae", "Neoaves", "Mirandornithes"],

    "Musophagiformes": ["Neornithes", "Neognathae", "Neoaves", "Otidimorphae"],
    "Otidiformes": ["Neornithes", "Neognathae", "Neoaves", "Otidimorphae"],
    "Cuculiformes": ["Neornithes", "Neognathae", "Neoaves", "Otidimorphae"],

    "Mesitornithiformes": ["Neornithes", "Neognathae", "Neoaves", "Columbimorphae"],
    "Pterocliformes": ["Neornithes", "Neognathae", "Neoaves", "Columbimorphae"],
    "Columbiformes": ["Neornithes", "Neognathae", "Neoaves", "Columbimorphae"],

    "Gaviiformes": ["Neornithes", "Neognathae", "Neoaves", "Aequornithes"],
    "Sphenisciformes": ["Neornithes", "Neognathae", "Neoaves", "Aequornithes"],
    "Procellariiformes": ["Neornithes", "Neognathae", "Neoaves", "Aequornithes"],
    "Ciconiiformes": ["Neornithes", "Neognathae", "Neoaves", "Aequornithes"],
    "Suliformes": ["Neornithes", "Neognathae", "Neoaves", "Aequornithes"],
    "Pelecaniformes": ["Neornithes", "Neognathae", "Neoaves", "Aequornithes"],

    "Caprimulgiformes": ["Neornithes", "Neognathae", "Neoaves", "Strisores"],
    "Steatornithiformes": ["Neornithes", "Neognathae", "Neoaves", "Strisores"],
    "Nyctibiiformes": ["Neornithes", "Neognathae", "Neoaves", "Strisores"],
    "Podargiformes": ["Neornithes", "Neognathae", "Neoaves", "Strisores"],
    "Aegotheliformes": ["Neornithes", "Neognathae", "Neoaves", "Strisores"],
    "Apodiformes": ["Neornithes", "Neognathae", "Neoaves", "Strisores"],

    "Phaethontiformes": ["Neornithes", "Neognathae", "Neoaves", "Eurypygimorphae"],
    "Eurypygiformes": ["Neornithes", "Neognathae", "Neoaves", "Eurypygimorphae"],

    "Accipitriformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],
    "Cathartiformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],
    "Strigiformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],
    "Coliiformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],
    "Leptosomiformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],
    "Trogoniformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],
    "Bucerotiformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],
    "Coraciiformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],
    "Galbuliformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],
    "Piciformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Afroaves"],

    "Cariamiformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Australaves"],
    "Falconiformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Australaves"],
    "Psittaciformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Australaves", "Psittacopasserae"],
    "Passeriformes": ["Neornithes", "Neognathae", "Neoaves", "Telluraves", "Australaves", "Psittacopasserae"],

    # These orders are retained at Neoaves level because deeper placement
    # is not encoded as settled in the MetaAves backbone.
    "Gruiformes": ["Neornithes", "Neognathae", "Neoaves"],
    "Charadriiformes": ["Neornithes", "Neognathae", "Neoaves"],
    "Opisthocomiformes": ["Neornithes", "Neognathae", "Neoaves"],
}


RANKS = ["class", "order", "family", "genus", "species"]


def clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def normalize_header(value):
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower()).strip("_")


def find_col(headers, candidates):
    normalized = {
        normalize_header(header): index
        for index, header in enumerate(headers)
        if header is not None
    }

    for candidate in candidates:
        key = normalize_header(candidate)
        if key in normalized:
            return normalized[key]

    for key, index in normalized.items():
        if any(normalize_header(candidate) in key for candidate in candidates):
            return index

    return None


def taxon_id(rank, name):
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_")
    return f"{rank}:{safe_name}"


def load_thai_names(path):
    """Load Thai eBird names from either wide or long-format common-name workbooks."""
    if not path:
        return {}

    wb = load_workbook(path, read_only=True, data_only=True)

    def header_kind(value):
        raw = clean(value) or ""
        normalized = normalize_header(raw)
        compact = normalized.replace("_", "")
        raw_lower = raw.lower()

        scientific = (
            "scientific" in compact
            or compact in {"scientificname", "scientificnames", "scientific_name"}
            or "ชื่อวิทยาศาสตร์" in raw
        )
        thai = (
            "thai" in raw_lower
            or compact in {"th", "thainame", "thainames"}
            or "ชื่อภาษาไทย" in raw
            or "ชื่อไทย" in raw
        )
        common = (
            "common" in compact
            or "name" == compact
            or compact in {"commonname", "commonnames", "birdname"}
        )
        language = (
            compact in {"language", "languagecode", "lang", "langcode", "locale", "languageid"}
            or "language" in compact
        )
        return scientific, thai, common, language

    # First support the wide format: one row contains Scientific Name and Thai.
    for ws in wb.worksheets:
        for row_index, row in enumerate(ws.iter_rows(values_only=True)):
            if row_index >= 100:
                break

            scientific_col = None
            thai_col = None
            common_col = None
            language_col = None

            for i, value in enumerate(row):
                scientific, thai, common, language = header_kind(value)
                if scientific and scientific_col is None:
                    scientific_col = i
                if thai and thai_col is None:
                    thai_col = i
                if common and common_col is None:
                    common_col = i
                if language and language_col is None:
                    language_col = i

            if scientific_col is None:
                continue

            # Wide spreadsheet with a dedicated Thai column.
            if thai_col is not None:
                thai_names = {}
                for data_row in ws.iter_rows(values_only=True, min_row=row_index + 2):
                    scientific = clean(data_row[scientific_col]) if scientific_col < len(data_row) else None
                    thai = clean(data_row[thai_col]) if thai_col < len(data_row) else None
                    if scientific and thai:
                        thai_names[scientific] = thai

                if thai_names:
                    print(f"Loaded {len(thai_names):,} Thai bird names from {ws.title!r}.")
                    return thai_names

    # Also support the long/matrix format used by eBird, where each row is a
    # scientific name + common name + language/code rather than a Thai column.
    for ws in wb.worksheets:
        for row_index, row in enumerate(ws.iter_rows(values_only=True)):
            if row_index >= 100:
                break

            scientific_col = common_col = language_col = None
            for i, value in enumerate(row):
                scientific, thai, common, language = header_kind(value)
                if scientific and scientific_col is None:
                    scientific_col = i
                if common and common_col is None:
                    common_col = i
                if language and language_col is None:
                    language_col = i

            if scientific_col is None or common_col is None:
                continue

            thai_names = {}

            # If the worksheet itself is clearly the Thai language sheet,
            # every common-name value is a Thai name.
            sheet_is_thai = "thai" in ws.title.lower() or "th" == ws.title.strip().lower()

            for data_row in ws.iter_rows(values_only=True, min_row=row_index + 2):
                scientific = clean(data_row[scientific_col]) if scientific_col < len(data_row) else None
                common = clean(data_row[common_col]) if common_col < len(data_row) else None
                language = clean(data_row[language_col]) if language_col is not None and language_col < len(data_row) else None

                if not scientific or not common:
                    continue

                is_thai = sheet_is_thai
                if language:
                    code = language.strip().lower().replace("-", "_")
                    is_thai = (
                        code in {"th", "th_th", "thai"}
                        or "thai" in code
                        or "ไทย" in language
                    )

                if is_thai:
                    thai_names[scientific] = common

            if thai_names:
                print(f"Loaded {len(thai_names):,} Thai bird names from {ws.title!r}.")
                return thai_names

    samples = []
    for ws in wb.worksheets:
        for row_index, row in enumerate(ws.iter_rows(values_only=True)):
            if row_index >= 8:
                break
            values = [str(v).strip() for v in row if v is not None]
            if values:
                samples.append(f"{ws.title!r}: {values}")

    raise SystemExit(
        "Could not find a usable scientific/Thai-name layout in the "
        "eBird common-names workbook. Header samples: "
        + " | ".join(samples[:12])
    )


def thai_name_for_bird(bird, thai_names):
    """Resolve a Thai name across taxonomy-name changes.

    AviList can move a species between genera while the Thai reference list
    still uses an older scientific combination. Prefer exact scientific-name
    matches, then match the species epithet only when the genus differs.
    This avoids losing established Thai names merely because taxonomy changed.
    """
    scientific = clean(bird.get("scientificName"))
    if not scientific:
        return None

    exact = thai_names.get(scientific)
    if exact:
        return exact

    parts = scientific.split()
    if len(parts) < 2:
        return None

    epithet = parts[1]
    candidates = [
        value for name, value in thai_names.items()
        if len(name.split()) >= 2 and name.split()[1] == epithet
    ]

    if len(candidates) == 1:
        return candidates[0]

    return None


def thai_name_for_bird(bird, thai_names):
    """Resolve Thai names even when AviList changes the scientific genus.

    Prefer an exact scientific-name match. If taxonomy has moved the species
    to another genus, fall back to a unique match on the species epithet.
    This prevents established Thai names from disappearing solely because
    the current taxonomy uses a different genus combination.
    """
    scientific = clean(bird.get("scientificName"))
    if not scientific:
        return None

    exact = thai_names.get(scientific)
    if exact:
        return exact

    parts = scientific.split()
    if len(parts) < 2:
        return None

    epithet = parts[1]
    candidates = [
        value for name, value in thai_names.items()
        if len(name.split()) >= 2 and name.split()[1] == epithet
    ]

    return candidates[0] if len(candidates) == 1 else None


def find_avilist_sheet(workbook):
    # Prefer the known current name, but normalize case/spacing so the
    # importer also works with harmless naming differences.
    preferred = "AviList v2025b extended"
    if preferred in workbook.sheetnames:
        return workbook[preferred]

    target = normalize_header(preferred)
    for sheet_name in workbook.sheetnames:
        normalized = normalize_header(sheet_name)
        if normalized == target:
            return workbook[sheet_name]

    # Fallback: find any sheet containing both "avilist" and "extended".
    for sheet_name in workbook.sheetnames:
        normalized = normalize_header(sheet_name)
        if "avilist" in normalized and "extended" in normalized:
            return workbook[sheet_name]

    available = ", ".join(repr(name) for name in workbook.sheetnames)
    raise SystemExit(
        "Could not find the AviList extended worksheet. "
        f"Available worksheets: {available}"
    )



def validate_import(birds, taxa):
    print("\n=== MetaAves Import Validation ===")

    missing_fields = {
        "commonName": [],
        "scientificName": [],
        "order": [],
        "family": [],
        "genus": []
    }

    for bird in birds:
        for field in missing_fields:
            if not bird.get(field):
                missing_fields[field].append(bird.get("scientificName") or bird.get("commonName") or "<unknown>")

    def duplicate_count(values):
        seen = set()
        duplicates = set()
        for value in values:
            if not value:
                continue
            if value in seen:
                duplicates.add(value)
            seen.add(value)
        return len(duplicates)

    duplicate_scientific = duplicate_count(
        bird.get("scientificName") for bird in birds
    )
    duplicate_common = duplicate_count(
        bird.get("commonName") for bird in birds
    )

    extinct_count = sum(1 for bird in birds if bird.get("isExtinct"))

    invalid_taxon_parents = [
        taxon_id_value
        for taxon_id_value, taxon in taxa.items()
        if taxon_id_value != "class:Aves"
        and taxon.get("parent") not in taxa
    ]

    root = taxa.get("class:Aves")
    invalid_root = (
        not root
        or root.get("rank") != "class"
        or root.get("name") != "Aves"
        or root.get("parent") is not None
    )

    print(f"Species imported:        {len(birds):,}")
    print(f"Extinct/possibly extinct:{extinct_count:,}")
    print(f"Taxonomy nodes:          {len(taxa):,}")
    print()
    print(f"Missing common names:    {len(missing_fields['commonName']):,}")
    print(f"Missing scientific names:{len(missing_fields['scientificName']):,}")
    print(f"Missing order:           {len(missing_fields['order']):,}")
    print(f"Missing family:          {len(missing_fields['family']):,}")
    print(f"Missing genus:           {len(missing_fields['genus']):,}")
    print()
    print(f"Duplicate scientific:    {duplicate_scientific:,}")
    print(f"Duplicate common names:  {duplicate_common:,}")
    print(f"Invalid taxon parents:   {len(invalid_taxon_parents):,}")
    print(f"Invalid Aves root:       {'YES' if invalid_root else 'NO'}")

    errors = (
        any(missing_fields.values())
        or duplicate_scientific > 0
        or invalid_taxon_parents
        or invalid_root
    )

    if errors:
        print("\nSTATUS: CHECK REQUIRED")
        for field, values in missing_fields.items():
            if values:
                print(f"  {field}: {values[:5]}")
        if invalid_taxon_parents:
            print(f"  Invalid parent examples: {invalid_taxon_parents[:5]}")
        return False

    print("\nSTATUS: PASS")
    return True


def main():
    if len(sys.argv) not in {2, 3}:
        raise SystemExit(
            "Usage: python scripts/import_avilist.py <AviList-extended.xlsx> "
            "[ebird-common-names.xlsx]"
        )

    xlsx = Path(sys.argv[1])
    ebird_names_path = Path(sys.argv[2]) if len(sys.argv) == 3 else None
    if not xlsx.exists():
        raise SystemExit(f"File not found: {xlsx}")

    thai_names = load_thai_names(ebird_names_path) if ebird_names_path else {}

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = find_avilist_sheet(wb)

    print(f"Using worksheet: {ws.title}")

    rows = ws.iter_rows(values_only=True)
    headers = next(rows)

    cols = {
        "rank": find_col(headers, ["Taxon_rank", "Rank"]),
        "english": find_col(headers, ["English_name_AviList", "English_name"]),
        "scientific": find_col(headers, ["Scientific_name"]),
        "order": find_col(headers, ["Order"]),
        "family": find_col(headers, ["Family"]),
        "genus": find_col(headers, ["Genus"]),
        "extinct": find_col(headers, ["Extinct_or_possibly_extinct"]),
        "range": find_col(headers, ["Range"]),
        "iucn": find_col(headers, ["IUCN_Red_List_Category"]),
    }

    missing = [
        key for key, value in cols.items()
        if value is None and key in {"rank", "english", "scientific"}
    ]
    if missing:
        raise SystemExit(
            "Could not find required AviList columns: "
            + ", ".join(missing)
        )

    birds = []
    clade_membership = {}
    unmapped_orders = set()

    taxa = {
        "class:Aves": {
            "id": "class:Aves",
            "rank": "class",
            "name": "Aves",
            "parent": None,
            "children": []
        }
    }

    for row in rows:
        rank = clean(row[cols["rank"]])

        if rank != "species":
            continue

        common = clean(row[cols["english"]])
        scientific = clean(row[cols["scientific"]])

        if not common or not scientific:
            continue

        # AviList's extended sheet may leave the Genus column blank even
        # though the scientific name is binomial. Use the explicit Genus
        # value when present; otherwise derive the genus from the first
        # nomenclatural token of the scientific name.
        genus = (
            clean(row[cols["genus"]])
            if cols["genus"] is not None
            else None
        )
        if not genus:
            genus = scientific.split()[0]

        bird = {
            "commonName": common,
            "scientificName": scientific,
            "thaiName": None,
            "isExtinct": False,
            "kingdom": "Animalia",
            "phylum": "Chordata",
            "class": "Aves",
            "order": clean(row[cols["order"]]) if cols["order"] is not None else None,
            "family": clean(row[cols["family"]]) if cols["family"] is not None else None,
            "genus": genus,
            "species": scientific,
            "habitat": None,
            "distribution": clean(row[cols["range"]]) if cols["range"] is not None else None,
            "diet": None,
            "behavior": None,
            "breeding": None,
            "conservation": clean(row[cols["iucn"]]) if cols["iucn"] is not None else None,
            "interestingFacts": [],
            "wikipediaTitle": common,
            "genusCharacteristics": []
        }

        extinct_value = clean(row[cols["extinct"]]) if cols["extinct"] is not None else None
        bird["thaiName"] = thai_name_for_bird(bird, thai_names)

        bird["isExtinct"] = bool(
            extinct_value
            and extinct_value.lower() in {
                "yes", "true", "extinct", "possibly extinct"
            }
        )

        birds.append(bird)

        order_name = bird.get("order")
        clade_path = CLADE_PATHS_BY_ORDER.get(order_name)
        if clade_path:
            clade_membership[scientific] = clade_path
        elif order_name:
            unmapped_orders.add(order_name)

        parent_id = "class:Aves"

        for tax_rank in RANKS[RANKS.index("class") + 1:]:
            value = bird.get(tax_rank)
            if not value:
                continue

            current_id = taxon_id(tax_rank, value)

            if current_id not in taxa:
                taxa[current_id] = {
                    "id": current_id,
                    "rank": tax_rank,
                    "name": value,
                    "parent": parent_id,
                    "children": []
                }

                if parent_id in taxa:
                    if current_id not in taxa[parent_id]["children"]:
                        taxa[parent_id]["children"].append(current_id)

            parent_id = current_id

        species_id = taxon_id("species", scientific)

        if species_id not in taxa:
            taxa[species_id] = {
                "id": species_id,
                "rank": "species",
                "name": scientific,
                "commonName": common,
                "parent": parent_id,
                "children": []
            }
            taxa[parent_id]["children"].append(species_id)

    out = Path("data")
    out.mkdir(exist_ok=True)

    (out / "clade_membership.generated.json").write_text(
        json.dumps(
            {
                "_meta": {
                    "version": 1,
                    "generatedBy": "scripts/import_avilist.py",
                    "source": "MetaAves broad phylogenetic backbone",
                    "policy": "Broad named clades only; contested deep Neoaves relationships are not forced."
                },
                "species": clade_membership
            },
            ensure_ascii=False,
            indent=2
        ),
        encoding="utf-8"
    )

    (out / "birds.generated.json").write_text(
        json.dumps(birds, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )

    taxonomy = {
        "_meta": {
            "version": 4,
            "masterSource": {
                "name": "AviList: The Global Avian Checklist",
                "version": "2025b"
            },
            "generatedBy": "scripts/import_avilist.py",
            "rankOrder": RANKS,
            "rootTaxa": ["class:Aves"],
            "nodeTypes": ["ranked_taxon", "species"],
            "taxonomyPolicy": "Classic MetaAves taxonomy uses only Class → Order → Family → Genus → Species. Named phylogenetic clades are maintained separately.",
            "clades": (
                "Maintained separately from ranked taxonomy. "
                "The game can insert clade nodes between ranked taxa."
            )
        }
    }
    taxonomy.update(taxa)

    (out / "taxonomy.generated.json").write_text(
        json.dumps(taxonomy, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )

    print(f"Imported {len(birds):,} species")
    print(f"Generated {len(taxa):,} ranked taxonomy nodes")
    print(f"Clade memberships: {len(clade_membership):,}")
    print(f"Orders without a clade mapping: {len(unmapped_orders):,}")
    if unmapped_orders:
        print("Unmapped orders:", ", ".join(sorted(unmapped_orders)))
    print("Wrote data/birds.generated.json")
    print("Wrote data/taxonomy.generated.json")
    print("Wrote data/clade_membership.generated.json")

    validate_import(birds, taxa)


if __name__ == "__main__":
    main()
